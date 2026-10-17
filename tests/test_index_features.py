"""Optional document features through real indexes and mocked network boundaries."""

import importlib
import asyncio
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from zipfile import ZipFile

import pytest

pytest.importorskip("llama_index.core")
pytest.importorskip("llama_index.llms.openai")
pytest.importorskip("langchain_openai")

from llama_index.core import Document, VectorStoreIndex
from llama_index.core.llms import MockLLM

from tests.test_indexing import integrations, mock_context  # noqa: F401


@pytest.mark.asyncio
async def test_context_index_and_chat_query_work_when_vocabulary_network_is_unavailable(
    monkeypatch, tmp_path
):
    import llama_index.core
    import tiktoken
    from requests.exceptions import ProxyError
    from llama_index.core.embeddings import MockEmbedding
    from llama_index.core.llms import LLMMetadata
    from collections import defaultdict

    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("ALLOWED_GUILDS", "1")
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("SHARE_DIR", str(tmp_path))
    usage = importlib.import_module("services.usage_service")
    monkeypatch.setattr(usage, "_tokenizer_unavailable", False)
    monkeypatch.setattr(llama_index.core, "global_tokenizer", None)
    attempts = []

    def unavailable(name):
        attempts.append(name)
        raise ProxyError("Vocabulary host unavailable")

    monkeypatch.setattr(tiktoken, "get_encoding", unavailable)
    service = importlib.import_module("services.indexing_service")
    module = importlib.import_module("models.index_model")
    context = service.create_index_context()

    class ChatMockLLM(MockLLM):
        @property
        def metadata(self):
            return LLMMetadata(
                context_window=4096,
                num_output=4,
                is_chat_model=True,
                model_name="offline-chat",
            )

    context.llm = ChatMockLLM(max_tokens=4, callback_manager=context.callback_manager)
    context.embed_model = MockEmbedding(
        embed_dim=8, callback_manager=context.callback_manager
    )
    index = VectorStoreIndex.from_documents(
        [Document(text="A document available during a vocabulary outage.")],
        **context.index_kwargs,
    )
    indexes = defaultdict(module.IndexData)
    indexes[1].queryable_index = index

    response = await asyncio.to_thread(
        module.get_and_query,
        1,
        indexes,
        "What is available?",
        "compact",
        1,
        1,
        context,
        False,
    )

    assert str(response) == "text text text text"
    assert len(response.source_nodes) == 1
    assert attempts == ["cl100k_base"]
    assert llama_index.core.global_tokenizer is None
    assert context.token_counter.total_embedding_token_count > 0
    assert context.token_counter.total_llm_token_count > 0


def test_optional_agent_memory_survives_unavailable_vocabulary(monkeypatch, tmp_path):
    import tiktoken
    from requests.exceptions import ProxyError
    from langchain_classic.memory import ConversationSummaryBufferMemory
    from services.agent_service import create_chat_llm
    from services import usage_service
    from models.openai_model import Models

    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setattr(usage_service, "_tokenizer_unavailable", False)
    attempts = []

    def unavailable(name):
        attempts.append(name)
        raise ProxyError("Vocabulary host unavailable")

    monkeypatch.setattr(tiktoken, "get_encoding", unavailable)
    llm = create_chat_llm(Models.DEFAULT)
    memory = ConversationSummaryBufferMemory(
        llm=llm, max_token_limit=16000, return_messages=True
    )
    memory.save_context({"input": "A source question"}, {"output": "A source answer"})
    memory.save_context(
        {"input": "A followup question"}, {"output": "A followup answer"}
    )
    assert len(memory.chat_memory.messages) == 4
    assert attempts == ["cl100k_base"]


class Response:
    def __init__(self, data=b"%PDF mocked", headers=None):
        self.headers = headers or {}
        self.data = data
        self.status = 200

    def raise_for_status(self):
        pass

    async def read(self):
        return self.data

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        return False


def session_factory(response, requests):
    class Session(Response):
        def __init__(self, **kwargs):
            requests.append(kwargs)

        def get(self, url, **kwargs):
            requests.append({"url": url, **kwargs})
            return response

    return Session


@pytest.mark.asyncio
@pytest.mark.parametrize("parse_fails", [False, True])
async def test_pdf_reader_cleans_download_on_success_and_parse_failure(
    integrations, monkeypatch, parse_fails
):
    files, requests = [], []
    monkeypatch.setattr(
        integrations.EnvService,
        "get_http_proxy",
        lambda url: "http://proxy.invalid:8080",
    )
    monkeypatch.setattr(
        integrations.aiohttp, "ClientSession", session_factory(Response(), requests)
    )

    class Reader:
        def __init__(self, input_files):
            files.extend(Path(path) for path in input_files)

        def load_data(self):
            assert files[0].read_bytes() == b"%PDF mocked"
            if parse_fails:
                raise ValueError("Invalid PDF")
            return [Document(text="Readable PDF document")]

    monkeypatch.setattr(integrations, "SimpleDirectoryReader", Reader)
    if parse_fails:
        with pytest.raises(ValueError, match="Invalid PDF"):
            await integrations.read_pdf_documents("https://example.com/report.pdf")
    else:
        documents = await integrations.read_pdf_documents(
            "https://example.com/report.pdf"
        )
        assert documents[0].metadata["URL"] == "https://example.com/report.pdf"
    assert files and not files[0].exists()
    assert requests[0]["timeout"].total == 30
    assert requests[1]["proxy"] == "http://proxy.invalid:8080"


@pytest.mark.asyncio
@pytest.mark.parametrize("content_type", [None, "application/pdf; charset=binary"])
async def test_web_index_accepts_missing_headers_and_pdf_parameters(
    integrations, monkeypatch, content_type
):
    module = importlib.import_module("models.index_model")
    context = mock_context(integrations)
    headers = {} if content_type is None else {"Content-Type": content_type}
    monkeypatch.setattr(
        module.aiohttp, "ClientSession", session_factory(Response(headers=headers), [])
    )
    handler = module.Index_handler(None, None)
    pdf = AsyncMock(return_value=[Document(text="PDF source document")])
    monkeypatch.setattr(handler, "index_pdf", pdf)
    monkeypatch.setattr(
        module,
        "read_web_documents",
        lambda url: [Document(text="HTML source document")],
    )

    index = await handler.index_webpage("https://example.com/report", context)

    source = " ".join(node.get_content() for node in index.docstore.docs.values())
    assert ("PDF" if content_type else "HTML") in source
    assert pdf.await_count == bool(content_type)


@pytest.mark.asyncio
async def test_discord_backup_keeps_embedding_model_for_later_reload(
    integrations, monkeypatch, tmp_path
):
    module = importlib.import_module("models.index_model")
    context = mock_context(integrations)
    context.embed_model.model_name = "text-embedding-3-small"
    monkeypatch.setattr(module, "create_index_context", lambda **kwargs: context)
    usage = SimpleNamespace(
        update_usage=AsyncMock(), get_price=AsyncMock(return_value=0.01)
    )
    handler = module.Index_handler(None, usage)
    handler.load_data = AsyncMock(
        return_value=[Document(text="Discord channel source")]
    )
    ctx = SimpleNamespace(
        guild=SimpleNamespace(
            id=12, name="Test Guild", text_channels=[SimpleNamespace(id=34)]
        ),
        respond=AsyncMock(),
    )

    await handler.backup_discord(ctx, user_api_key="personal-key", message_limit=100)

    metadata = list((tmp_path / "indexes" / "12").glob("*/gptdiscord-index.json"))
    assert len(metadata) == 1
    assert '"text-embedding-3-small"' in metadata[0].read_text()
    assert usage.update_usage.await_args.args[1] == "text-embedding-3-small"
    handler.load_data.assert_awaited_once_with(
        channel_ids=[34], limit=100, oldest_first=False
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("deep", [False, True])
async def test_saved_indexes_compose_with_current_llama_index_apis(
    integrations, monkeypatch, tmp_path, deep
):
    module = importlib.import_module("models.index_model")

    class NamedMockLLM(MockLLM):
        model: str = "gpt-6-luna"

    context = mock_context(integrations)
    context.llm = NamedMockLLM(max_tokens=4, callback_manager=context.callback_manager)
    context.embed_model.model_name = "text-embedding-3-small"
    indexes = [
        VectorStoreIndex.from_documents([Document(text=text)], **context.index_kwargs)
        for text in ("First indexed source", "Second indexed source")
    ]
    usage = SimpleNamespace(
        update_usage=AsyncMock(), get_price=AsyncMock(return_value=0.01)
    )
    handler = module.Index_handler(None, usage)
    monkeypatch.setattr(module.EnvService, "find_shared_file", lambda name: Path(name))
    monkeypatch.setattr(
        handler,
        "index_load_file",
        lambda path, ctx: indexes[0 if str(path).endswith("first") else 1],
    )

    price = await handler.compose_indexes(
        42, ["first", "second"], "combined", deep, context
    )

    assert price > 0
    assert handler.index_storage[42].queryable_index is not None
    assert (tmp_path / "indexes" / "42" / "combined" / "docstore.json").exists()
    assert usage.update_usage.await_count == (2 if deep else 1)


def write_document(path, suffix):
    """Use standard formats and the installed production reader dependencies."""
    if suffix == ".docx":
        with ZipFile(path, "w") as archive:
            archive.writestr(
                "word/document.xml",
                (
                    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    "<w:body><w:p><w:r><w:t>Codename FENNEL</w:t></w:r></w:p></w:body></w:document>"
                ),
            )
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = (
            "Codename FENNEL"
        )
        presentation.save(path)
    elif suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Property", "Value"])
        workbook.active.append(["Codename", "FENNEL"])
        workbook.save(path)
        workbook.close()
    elif suffix == ".epub":
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("offline-parser-fixture")
        book.set_title("Parser fixture")
        book.set_language("en")
        chapter = epub.EpubHtml(title="Codename", file_name="chapter.xhtml", lang="en")
        chapter.content = "<html><body><p>Codename FENNEL</p></body></html>"
        book.add_item(chapter)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.toc = (chapter,)
        book.spine = ["nav", chapter]
        epub.write_epub(str(path), book)
    elif suffix == ".pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Codename FENNEL) Tj ET")
        page[NameObject("/Contents")] = stream
        with path.open("wb") as output:
            writer.write(output)
    else:
        path.write_text(
            {
                ".md": "# Codename\n\nFENNEL",
                ".csv": "property,value\ncodename,FENNEL\n",
                ".rtf": r"{\rtf1\ansi Codename FENNEL}",
            }.get(suffix, "Codename FENNEL")
        )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "suffix",
    [".txt", ".md", ".csv", ".pdf", ".docx", ".pptx", ".xlsx", ".epub", ".rtf"],
)
async def test_advertised_document_formats_extract_text_into_real_index(
    integrations, tmp_path, suffix
):
    module = importlib.import_module("models.index_model")
    path = tmp_path / f"fixture{suffix}"
    write_document(path, suffix)
    context = mock_context(integrations)
    handler = module.Index_handler(None, None)

    index = handler.index_file(path, context, suffix)

    assert any("FENNEL" in node.get_content() for node in index.docstore.docs.values())
    assert context.token_counter.total_embedding_token_count > 0


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", [".png", ".mp3", ".mp4", ".doc", ".ppt", ".xls"])
async def test_unsupported_media_and_legacy_binary_formats_never_embed(
    integrations, tmp_path, suffix
):
    module = importlib.import_module("models.index_model")
    context = mock_context(integrations)
    path = tmp_path / f"fixture{suffix}"
    path.write_bytes(b"Not document text")
    handler = module.Index_handler(None, None)

    with pytest.raises(ValueError, match="unavailable|binary format"):
        handler.index_file(path, context, suffix)

    assert context.token_counter.total_embedding_token_count == 0


def index_session(integrations, monkeypatch):
    module = importlib.import_module("cogs.index_service_cog")
    monkeypatch.setattr(module, "USER_INPUT_API_KEYS", False)
    monkeypatch.setattr(module, "PRE_MODERATE", False)
    cog = module.IndexService(
        SimpleNamespace(user=SimpleNamespace(id=99)), None, asyncio.Queue()
    )
    cog.index_handler.index_chat_chains[5] = SimpleNamespace(owner_id=42)
    cog.index_handler.execute_index_chat_message = AsyncMock(
        return_value="Agent answer"
    )
    message = SimpleNamespace(
        author=SimpleNamespace(id=42, bot=False),
        guild=SimpleNamespace(id=1),
        type=module.discord.MessageType.reply,
        content="A document question",
        attachments=[],
        channel=SimpleNamespace(id=5, trigger_typing=AsyncMock()),
        reply=AsyncMock(),
        delete=AsyncMock(),
    )
    return module, cog, message


@pytest.mark.asyncio
@pytest.mark.parametrize("mime", [None, "application/octet-stream", "APPLICATION/PDF"])
async def test_upload_type_detection_uses_case_insensitive_extension_fallback(
    integrations, mime
):
    module = importlib.import_module("models.index_model")
    handler = module.Index_handler(None, None)
    assert handler.get_file_suffix(mime, "REPORT.PDF") == ".pdf"
    assert handler.get_file_suffix(None, "REPORT.DOCX") == ".docx"


@pytest.mark.asyncio
async def test_index_session_reserves_before_moderation_await(
    integrations, monkeypatch
):
    module, cog, message = index_session(integrations, monkeypatch)
    monkeypatch.setattr(module, "PRE_MODERATE", True)
    entered, finish = asyncio.Event(), asyncio.Event()

    async def moderate(*args):
        entered.set()
        await finish.wait()
        return False

    moderation = AsyncMock(side_effect=moderate)
    monkeypatch.setattr(module.Moderation, "simple_moderate_and_respond", moderation)
    first = asyncio.create_task(cog.on_message(message))
    await asyncio.wait_for(entered.wait(), 1)
    await cog.on_message(message)
    finish.set()
    await first

    moderation.assert_awaited_once()
    cog.index_handler.execute_index_chat_message.assert_awaited_once()
    assert not cog.thread_awaiting_responses


@pytest.mark.asyncio
@pytest.mark.parametrize("outcome", [None, "reply-fails"])
async def test_index_session_clears_busy_after_empty_output_or_reply_failure(
    integrations, monkeypatch, outcome
):
    _, cog, message = index_session(integrations, monkeypatch)
    if outcome is None:
        cog.index_handler.execute_index_chat_message.return_value = None
        await cog.on_message(message)
    else:
        message.reply.side_effect = OSError("Discord write failed")
        with pytest.raises(OSError):
            await cog.on_message(message)
    assert not cog.thread_awaiting_responses


@pytest.mark.asyncio
async def test_index_session_awaits_all_long_answer_pages_and_ignores_other_bots(
    integrations, monkeypatch
):
    _, cog, message = index_session(integrations, monkeypatch)
    cog.index_handler.execute_index_chat_message.return_value = "Agent answer. " * 600
    page = SimpleNamespace()
    page.reply = AsyncMock(return_value=page)
    message.reply.return_value = page
    await cog.on_message(message)
    assert page.reply.await_count >= 2
    assert not cog.thread_awaiting_responses

    cog.index_handler.execute_index_chat_message.reset_mock()
    message.author.bot = True
    await cog.on_message(message)
    cog.index_handler.execute_index_chat_message.assert_not_awaited()
